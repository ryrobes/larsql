"""
Lazy ATTACH for configured external data sources.

Goal: make queries like `SELECT * FROM prod_db.public.users` "just work" by
attaching the referenced connection from `sql_connections/*.yaml` on demand.

Design principles:
- Best-effort: never fail a query just because an auto-attach failed.
- Low false-positives: default detection focuses on relation references (FROM/JOIN/UPDATE/INTO).
- No credential persistence: auto-attached connections are NOT written to _rvbbit_attachments.
"""

from __future__ import annotations

import logging
import os
import re
import shutil
import tempfile
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from .config import SqlConnectionConfig, resolve_google_credentials
from .connector import sanitize_name

log = logging.getLogger(__name__)


@dataclass(frozen=True)
class _Token:
    typ: str  # word, ident, punct, string, other
    text: str


def _tokenize(sql: str) -> List[_Token]:
    """
    Very small SQL tokenizer for lazy attach detection.

    - Skips comments and string literals (single-quoted and dollar-quoted).
    - Treats double-quoted identifiers as ident tokens.
    - Emits punctuation tokens for '.', ',', '(', ')', ';'.
    """
    tokens: List[_Token] = []
    i = 0
    n = len(sql)

    def emit(typ: str, start: int, end: int) -> None:
        if end > start:
            tokens.append(_Token(typ, sql[start:end]))

    while i < n:
        ch = sql[i]

        # Whitespace
        if ch.isspace():
            i += 1
            continue

        # Line comment
        if ch == "-" and i + 1 < n and sql[i + 1] == "-":
            i += 2
            while i < n and sql[i] != "\n":
                i += 1
            continue

        # Block comment
        if ch == "/" and i + 1 < n and sql[i + 1] == "*":
            i += 2
            while i + 1 < n and not (sql[i] == "*" and sql[i + 1] == "/"):
                i += 1
            i = min(n, i + 2)
            continue

        # Single-quoted string
        if ch == "'":
            i += 1
            while i < n:
                if sql[i] == "'":
                    if i + 1 < n and sql[i + 1] == "'":  # escaped ''
                        i += 2
                        continue
                    i += 1
                    break
                i += 1
            continue

        # Dollar-quoted string ($$...$$ or $tag$...$tag$)
        if ch == "$":
            m = re.match(r"\$[A-Za-z0-9_]*\$", sql[i:])
            if m:
                tag = m.group(0)
                i += len(tag)
                end = sql.find(tag, i)
                if end == -1:
                    # Unclosed; treat rest as string
                    break
                i = end + len(tag)
                continue

        # Double-quoted identifier
        if ch == '"':
            start = i + 1
            i += 1
            buf = []
            while i < n:
                if sql[i] == '"':
                    if i + 1 < n and sql[i + 1] == '"':  # escaped ""
                        buf.append('"')
                        i += 2
                        continue
                    i += 1
                    break
                buf.append(sql[i])
                i += 1
            tokens.append(_Token("ident", "".join(buf)))
            continue

        # Punctuation
        if ch in (".", ",", "(", ")", ";"):
            tokens.append(_Token("punct", ch))
            i += 1
            continue

        # Word/identifier
        if ch.isalpha() or ch == "_":
            start = i
            i += 1
            while i < n and (sql[i].isalnum() or sql[i] == "_"):
                i += 1
            emit("word", start, i)
            continue

        # Fallback
        tokens.append(_Token("other", ch))
        i += 1

    return tokens


def _parse_qualified_name(tokens: Sequence[_Token], start: int) -> Tuple[List[str], int]:
    """
    Parse `ident(.ident)+` starting at `start`.

    Returns: (parts, next_index)
    """
    if start >= len(tokens) or tokens[start].typ not in ("word", "ident"):
        return [], start

    parts = [tokens[start].text]
    i = start + 1
    while i + 1 < len(tokens) and tokens[i].typ == "punct" and tokens[i].text == "." and tokens[i + 1].typ in ("word", "ident"):
        parts.append(tokens[i + 1].text)
        i += 2
    return parts, i


_RELATION_KEYWORDS = {
    "FROM",
    "JOIN",
    "UPDATE",
    "INTO",
    "DELETE",
}

_RELATION_STOP_KEYWORDS = {
    "WHERE",
    "GROUP",
    "ORDER",
    "HAVING",
    "LIMIT",
    "QUALIFY",
    "WINDOW",
    "UNION",
    "EXCEPT",
    "INTERSECT",
    "RETURNING",
}

_RELATION_PREFIX_KEYWORDS = {"LATERAL", "ONLY"}


def extract_relation_qualified_names(sql: str) -> List[List[str]]:
    """
    Extract qualified relation names from common relation positions:
    - FROM <rel>
    - JOIN <rel>
    - UPDATE <rel>
    - INSERT INTO <rel>
    - DELETE FROM <rel>

    This is intentionally conservative to avoid capturing `t.col` expressions.
    """
    tokens = _tokenize(sql)
    out: List[List[str]] = []

    i = 0
    while i < len(tokens):
        tok = tokens[i]
        if tok.typ != "word":
            i += 1
            continue

        kw = tok.text.upper()
        if kw not in _RELATION_KEYWORDS:
            i += 1
            continue

        # DELETE FROM <rel>
        if kw == "DELETE":
            if i + 1 < len(tokens) and tokens[i + 1].typ == "word" and tokens[i + 1].text.upper() == "FROM":
                i += 1  # treat as FROM
                kw = "FROM"
            else:
                i += 1
                continue

        # INTO is only interesting for INSERT INTO (we don't enforce INSERT)
        i += 1

        # Skip optional modifiers like LATERAL/ONLY
        while i < len(tokens) and tokens[i].typ == "word" and tokens[i].text.upper() in _RELATION_PREFIX_KEYWORDS:
            i += 1

        # Subquery or derived table: FROM (SELECT ...)
        if i < len(tokens) and tokens[i].typ == "punct" and tokens[i].text == "(":
            # don't attempt to skip; inner FROM/JOIN will be caught by the main scan
            i += 1
            continue

        parts, next_i = _parse_qualified_name(tokens, i)
        if len(parts) >= 2:
            out.append(parts)

        i = max(i + 1, next_i)

    return out


def extract_any_dotted_prefixes(sql: str) -> Set[str]:
    """
    Aggressive fallback: extract the left-most identifier of any dotted path `x.y`.
    Intended for retry-on-error when conservative extraction misses.
    """
    tokens = _tokenize(sql)
    prefixes: Set[str] = set()

    i = 0
    while i + 2 < len(tokens):
        if tokens[i].typ in ("word", "ident") and tokens[i + 1].typ == "punct" and tokens[i + 1].text == "." and tokens[i + 2].typ in ("word", "ident"):
            # Only collect the LEFT-most segment in a chain like a.b.c
            if i == 0 or not (tokens[i - 1].typ == "punct" and tokens[i - 1].text == "."):
                prefixes.add(tokens[i].text)
            i += 1
            continue
        i += 1

    return prefixes


def _quote_ident(name: str) -> str:
    return f'"{name.replace(chr(34), chr(34) * 2)}"'


def _escape_single_quotes(value: str) -> str:
    return value.replace("'", "''")


class LazyAttachManager:
    """
    Per-connection lazy attach manager.

    Keeps small caches (failed attachments, folder file maps) to avoid repeatedly
    retrying broken connections on every query.
    """

    def __init__(self, duckdb_conn, sql_connections: Dict[str, SqlConnectionConfig]):
        self._conn = duckdb_conn
        self._configs = dict(sql_connections)

        # Some "connections" are folder-based and generate multiple attach targets
        self._duckdb_folder_configs = [
            cfg for cfg in self._configs.values() if cfg.type == "duckdb_folder" and cfg.folder_path
        ]

        self._failed_configs: Set[str] = set()
        self._duckdb_file_maps: Dict[str, Dict[str, Path]] = {}  # folder_path -> {db_name: file_path}
        self._csv_file_maps: Dict[str, Dict[str, Path]] = {}  # connection_name -> {table_name: file_path}
        self._clickhouse_tables: Dict[str, Set[str]] = {}  # connection_name -> {table_names already materialized}

    def ensure_for_query(self, sql: str, *, aggressive: bool = False) -> None:
        """
        Ensure any referenced configured sources are attached/materialized.

        If aggressive=True, falls back to scanning any dotted prefixes in addition
        to conservative FROM/JOIN/etc extraction.
        """
        if not _lazy_attach_enabled():
            return

        qualified = extract_relation_qualified_names(sql)

        dotted_prefixes: Set[str] = set()
        if aggressive:
            dotted_prefixes = extract_any_dotted_prefixes(sql)

        needed_catalogs: Set[str] = set()
        needed_csv_tables: Dict[str, Set[str]] = {}
        needed_clickhouse_tables: Dict[str, Set[str]] = {}

        # From relation-qualified names
        for parts in qualified:
            if not parts:
                continue
            prefix = parts[0]
            if prefix in self._configs:
                cfg = self._configs[prefix]
                if cfg.type == "csv_folder" and len(parts) >= 2:
                    needed_csv_tables.setdefault(prefix, set()).add(parts[1])
                elif cfg.type == "clickhouse" and len(parts) >= 2:
                    needed_clickhouse_tables.setdefault(prefix, set()).add(parts[1])
                else:
                    needed_catalogs.add(prefix)
            else:
                # Might be a DuckDB file in a duckdb_folder connection
                needed_catalogs.add(prefix)

        # From aggressive dotted prefixes
        for prefix in dotted_prefixes:
            if prefix in self._configs:
                cfg = self._configs[prefix]
                if cfg.type == "csv_folder":
                    # Not enough info to know which table; skip
                    continue
                if cfg.type == "clickhouse":
                    # Not enough info to know which table; skip
                    continue
                needed_catalogs.add(prefix)
            else:
                needed_catalogs.add(prefix)

        # Attach catalogs (postgres/mysql/sqlite or duckdb files)
        for catalog in sorted(needed_catalogs):
            self._ensure_catalog_attached(catalog)

        # Ensure CSV schema/tables (view/table) exist
        for schema_name, tables in needed_csv_tables.items():
            self._ensure_csv_tables(schema_name, tables)

        # Ensure ClickHouse tables are materialized
        for schema_name, tables in needed_clickhouse_tables.items():
            self._ensure_clickhouse_tables(schema_name, tables)

    # ---------------------------------------------------------------------
    # Attachment helpers
    # ---------------------------------------------------------------------

    def _attached_catalogs(self) -> Set[str]:
        try:
            rows = self._conn.execute(
                """
                SELECT database_name
                FROM duckdb_databases()
                WHERE NOT internal
                """
            ).fetchall()
            return {r[0] for r in rows}
        except Exception:
            return set()

    def _ensure_catalog_attached(self, catalog: str) -> None:
        # Already attached?
        if catalog in self._attached_catalogs():
            return

        # Known configured connection?
        cfg = self._configs.get(catalog)
        if cfg:
            if cfg.connection_name in self._failed_configs:
                return
            try:
                self._attach_config(cfg)
            except Exception as e:
                self._failed_configs.add(cfg.connection_name)
                log.debug("[lazy_attach] Failed attaching %s: %s", cfg.connection_name, e)
            return

        # Otherwise: see if it matches a duckdb_folder file name
        try:
            attached = self._attach_duckdb_file_if_present(catalog)
            if attached:
                return
        except Exception as e:
            log.debug("[lazy_attach] Failed attaching duckdb file %s: %s", catalog, e)

    def _attach_config(self, cfg: SqlConnectionConfig) -> None:
        if cfg.type == "postgres":
            self._attach_postgres(cfg)
            return
        if cfg.type == "mysql":
            self._attach_mysql(cfg)
            return
        if cfg.type == "sqlite":
            self._attach_sqlite(cfg)
            return
        # Phase 1: Cloud databases
        if cfg.type == "bigquery":
            self._attach_bigquery(cfg)
            return
        if cfg.type == "snowflake":
            self._attach_snowflake(cfg)
            return
        if cfg.type == "motherduck":
            self._attach_motherduck(cfg)
            return
        # Phase 2: Remote filesystems
        if cfg.type == "s3":
            self._attach_s3(cfg)
            return
        if cfg.type == "azure":
            self._attach_azure(cfg)
            return
        if cfg.type == "gcs":
            self._attach_gcs(cfg)
            return
        if cfg.type == "http":
            self._attach_http(cfg)
            return
        # Phase 3: Lakehouse formats
        if cfg.type == "delta":
            self._attach_delta(cfg)
            return
        if cfg.type == "iceberg":
            self._attach_iceberg(cfg)
            return
        # Phase 4: Scanner functions
        if cfg.type == "odbc":
            self._attach_odbc(cfg)
            return
        if cfg.type == "gsheets":
            self._attach_gsheets(cfg)
            return
        if cfg.type == "excel":
            self._attach_excel(cfg)
            return
        # Phase 5: Hybrid/materialization (handled separately due to table-based lazy attach)
        if cfg.type == "mongodb":
            self._attach_mongodb(cfg)
            return
        if cfg.type == "cassandra":
            self._attach_cassandra(cfg)
            return
        if cfg.type == "clickhouse":
            self._attach_clickhouse(cfg)
            return
        # csv_folder is handled via schema/table materialization (not catalog attach)
        if cfg.type == "csv_folder":
            return
        if cfg.type == "duckdb_folder":
            # Individual duckdb files are attached on demand by name.
            return
        raise ValueError(f"Unsupported connection type for lazy attach: {cfg.type}")

    def _attach_postgres(self, cfg: SqlConnectionConfig) -> None:
        # Best-effort extension setup
        try:
            self._conn.execute("INSTALL postgres;")
        except Exception:
            pass
        self._conn.execute("LOAD postgres;")

        if not (cfg.database and cfg.host and cfg.port and cfg.user):
            raise ValueError(f"postgres config missing required fields for {cfg.connection_name}")

        conn_str = f"dbname={cfg.database} host={cfg.host} port={cfg.port} user={cfg.user}"
        if cfg.password:
            conn_str += f" password={cfg.password}"

        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(
            f"ATTACH '{_escape_single_quotes(conn_str)}' AS {alias} (TYPE postgres);"
        )

    def _attach_mysql(self, cfg: SqlConnectionConfig) -> None:
        try:
            self._conn.execute("INSTALL mysql;")
        except Exception:
            pass
        self._conn.execute("LOAD mysql;")

        if not (cfg.database and cfg.host and cfg.port and cfg.user):
            raise ValueError(f"mysql config missing required fields for {cfg.connection_name}")

        conn_str = f"host={cfg.host} port={cfg.port} database={cfg.database} user={cfg.user}"
        if cfg.password:
            conn_str += f" password={cfg.password}"

        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(
            f"ATTACH '{_escape_single_quotes(conn_str)}' AS {alias} (TYPE mysql);"
        )

    def _attach_sqlite(self, cfg: SqlConnectionConfig) -> None:
        if not cfg.database:
            raise ValueError(f"sqlite config missing database path for {cfg.connection_name}")
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(
            f"ATTACH '{_escape_single_quotes(cfg.database)}' AS {alias} (TYPE sqlite);"
        )

    # -------------------------------------------------------------------------
    # Phase 1: Cloud Database Connectors
    # -------------------------------------------------------------------------

    def _attach_bigquery(self, cfg: SqlConnectionConfig) -> None:
        """Attach BigQuery database via DuckDB extension."""
        try:
            self._conn.execute("INSTALL bigquery FROM community;")
        except Exception:
            pass
        self._conn.execute("LOAD bigquery;")

        if not cfg.project_id:
            raise ValueError(f"bigquery config missing project_id for {cfg.connection_name}")

        # Handle credentials - BigQuery extension uses GOOGLE_APPLICATION_CREDENTIALS
        # Use resolver which handles both file paths and JSON strings
        if cfg.credentials_env:
            creds_path = resolve_google_credentials(cfg.credentials_env)
            if creds_path:
                os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = creds_path
        else:
            # Fallback to default GOOGLE_APPLICATION_CREDENTIALS
            creds_path = resolve_google_credentials()
            if creds_path:
                os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = creds_path

        conn_str = f"project={cfg.project_id}"
        read_only_clause = ", READ_ONLY" if cfg.read_only else ""
        alias = _quote_ident(cfg.connection_name)

        self._conn.execute(
            f"ATTACH '{_escape_single_quotes(conn_str)}' AS {alias} (TYPE bigquery{read_only_clause});"
        )

    def _attach_snowflake(self, cfg: SqlConnectionConfig) -> None:
        """Attach Snowflake database via DuckDB extension."""
        try:
            self._conn.execute("INSTALL snowflake;")
        except Exception:
            pass
        self._conn.execute("LOAD snowflake;")

        if not cfg.account or not cfg.user:
            raise ValueError(f"snowflake config missing account/user for {cfg.connection_name}")

        conn_parts = [f"account={cfg.account}", f"user={cfg.user}"]
        if cfg.password:
            conn_parts.append(f"password={cfg.password}")
        if cfg.database:
            conn_parts.append(f"database={cfg.database}")
        if cfg.warehouse:
            conn_parts.append(f"warehouse={cfg.warehouse}")
        if cfg.role:
            conn_parts.append(f"role={cfg.role}")

        conn_str = ";".join(conn_parts)
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(
            f"ATTACH '{_escape_single_quotes(conn_str)}' AS {alias} (TYPE snowflake);"
        )

    def _attach_motherduck(self, cfg: SqlConnectionConfig) -> None:
        """Attach Motherduck database."""
        if not cfg.database:
            raise ValueError(f"motherduck config missing database for {cfg.connection_name}")

        # Set token if provided
        if cfg.motherduck_token_env:
            token = os.getenv(cfg.motherduck_token_env)
            if token:
                self._conn.execute(f"SET motherduck_token='{_escape_single_quotes(token)}';")

        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(
            f"ATTACH 'md:{_escape_single_quotes(cfg.database)}' AS {alias};"
        )

    # -------------------------------------------------------------------------
    # Phase 2: Remote Filesystem Connectors
    # -------------------------------------------------------------------------

    def _attach_s3(self, cfg: SqlConnectionConfig) -> None:
        """Attach S3 bucket as a DuckDB schema."""
        try:
            self._conn.execute("INSTALL httpfs;")
        except Exception:
            pass
        self._conn.execute("LOAD httpfs;")

        # Configure S3 credentials
        region = cfg.region or "us-east-1"
        self._conn.execute(f"SET s3_region='{region}';")

        if cfg.access_key_env:
            access_key = os.getenv(cfg.access_key_env, "")
            if access_key:
                self._conn.execute(f"SET s3_access_key_id='{access_key}';")

        if cfg.secret_key_env:
            secret_key = os.getenv(cfg.secret_key_env, "")
            if secret_key:
                self._conn.execute(f"SET s3_secret_access_key='{secret_key}';")

        # Support S3-compatible endpoints (MinIO, R2, etc.)
        if cfg.endpoint_url:
            endpoint = cfg.endpoint_url.replace('http://', '').replace('https://', '')
            self._conn.execute(f"SET s3_endpoint='{endpoint}';")
            self._conn.execute("SET s3_url_style='path';")
            # Disable SSL for localhost endpoints
            if 'localhost' in cfg.endpoint_url or '127.0.0.1' in cfg.endpoint_url:
                self._conn.execute("SET s3_use_ssl=false;")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Build S3 path and create views for files
        prefix = cfg.prefix.rstrip('/') if cfg.prefix else ""
        s3_base = f"s3://{cfg.bucket}/{prefix}" if prefix else f"s3://{cfg.bucket}"
        pattern = cfg.file_pattern or "*.parquet"

        try:
            glob_path = f"{s3_base.rstrip('/')}/{pattern}"
            files = self._conn.execute(f"SELECT * FROM glob('{glob_path}')").fetchall()
            for row in files:
                file_path = row[0]
                table_name = sanitize_name(Path(file_path).name)
                try:
                    if file_path.endswith('.parquet'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_parquet('{file_path}')")
                    elif file_path.endswith('.csv'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_csv_auto('{file_path}')")
                except Exception as e:
                    log.debug("[lazy_attach] Failed creating S3 view %s: %s", table_name, e)
        except Exception as e:
            log.debug("[lazy_attach] S3 glob failed for %s: %s", cfg.connection_name, e)

    def _attach_azure(self, cfg: SqlConnectionConfig) -> None:
        """Attach Azure Blob Storage as a DuckDB schema."""
        try:
            self._conn.execute("INSTALL azure;")
        except Exception:
            pass
        self._conn.execute("LOAD azure;")

        # Configure Azure credentials
        if cfg.connection_string_env:
            conn_str = os.getenv(cfg.connection_string_env, "")
            if conn_str:
                self._conn.execute(f"SET azure_storage_connection_string='{conn_str}';")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Build Azure path and create views
        prefix = cfg.prefix.rstrip('/') if cfg.prefix else ""
        azure_base = f"azure://{cfg.bucket}/{prefix}" if prefix else f"azure://{cfg.bucket}"
        pattern = cfg.file_pattern or "*.parquet"

        try:
            glob_path = f"{azure_base.rstrip('/')}/{pattern}"
            files = self._conn.execute(f"SELECT * FROM glob('{glob_path}')").fetchall()
            for row in files:
                file_path = row[0]
                table_name = sanitize_name(Path(file_path).name)
                try:
                    if file_path.endswith('.parquet'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_parquet('{file_path}')")
                    elif file_path.endswith('.csv'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_csv_auto('{file_path}')")
                except Exception as e:
                    log.debug("[lazy_attach] Failed creating Azure view %s: %s", table_name, e)
        except Exception as e:
            log.debug("[lazy_attach] Azure glob failed for %s: %s", cfg.connection_name, e)

    def _attach_gcs(self, cfg: SqlConnectionConfig) -> None:
        """Attach Google Cloud Storage as a DuckDB schema."""
        try:
            self._conn.execute("INSTALL httpfs;")
        except Exception:
            pass
        self._conn.execute("LOAD httpfs;")

        # Configure GCS credentials - use resolver which handles both file paths and JSON strings
        if cfg.credentials_env:
            creds_path = resolve_google_credentials(cfg.credentials_env)
            if creds_path:
                os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = creds_path
        else:
            # Fallback to default GOOGLE_APPLICATION_CREDENTIALS
            creds_path = resolve_google_credentials()
            if creds_path:
                os.environ['GOOGLE_APPLICATION_CREDENTIALS'] = creds_path

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Build GCS path and create views
        prefix = cfg.prefix.rstrip('/') if cfg.prefix else ""
        gcs_base = f"gcs://{cfg.bucket}/{prefix}" if prefix else f"gcs://{cfg.bucket}"
        pattern = cfg.file_pattern or "*.parquet"

        try:
            glob_path = f"{gcs_base.rstrip('/')}/{pattern}"
            files = self._conn.execute(f"SELECT * FROM glob('{glob_path}')").fetchall()
            for row in files:
                file_path = row[0]
                table_name = sanitize_name(Path(file_path).name)
                try:
                    if file_path.endswith('.parquet'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_parquet('{file_path}')")
                    elif file_path.endswith('.csv'):
                        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_csv_auto('{file_path}')")
                except Exception as e:
                    log.debug("[lazy_attach] Failed creating GCS view %s: %s", table_name, e)
        except Exception as e:
            log.debug("[lazy_attach] GCS glob failed for %s: %s", cfg.connection_name, e)

    def _attach_http(self, cfg: SqlConnectionConfig) -> None:
        """Attach HTTP-accessible file as a DuckDB schema."""
        try:
            self._conn.execute("INSTALL httpfs;")
        except Exception:
            pass
        self._conn.execute("LOAD httpfs;")

        if not cfg.folder_path:
            raise ValueError(f"http config missing folder_path for {cfg.connection_name}")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Create view for HTTP file
        base_url = cfg.folder_path
        table_name = sanitize_name(Path(base_url).name) or "data"

        try:
            if base_url.endswith('.parquet'):
                self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_parquet('{base_url}')")
            elif base_url.endswith('.csv'):
                self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_csv_auto('{base_url}')")
            elif base_url.endswith('.json') or base_url.endswith('.jsonl'):
                self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_json_auto('{base_url}')")
            else:
                self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_parquet('{base_url}')")
        except Exception as e:
            log.debug("[lazy_attach] HTTP attach failed for %s: %s", cfg.connection_name, e)

    # -------------------------------------------------------------------------
    # Phase 3: Lakehouse Format Connectors
    # -------------------------------------------------------------------------

    def _attach_delta(self, cfg: SqlConnectionConfig) -> None:
        """Attach Delta Lake table."""
        try:
            self._conn.execute("INSTALL delta;")
        except Exception:
            pass
        self._conn.execute("LOAD delta;")

        # Configure S3 if needed
        if cfg.table_path and cfg.table_path.startswith('s3://'):
            try:
                self._conn.execute("INSTALL httpfs;")
            except Exception:
                pass
            self._conn.execute("LOAD httpfs;")

            region = cfg.region or "us-east-1"
            self._conn.execute(f"SET s3_region='{region}';")

            if cfg.access_key_env:
                access_key = os.getenv(cfg.access_key_env, "")
                if access_key:
                    self._conn.execute(f"SET s3_access_key_id='{access_key}';")

            if cfg.secret_key_env:
                secret_key = os.getenv(cfg.secret_key_env, "")
                if secret_key:
                    self._conn.execute(f"SET s3_secret_access_key='{secret_key}';")

        if not cfg.table_path:
            raise ValueError(f"delta config missing table_path for {cfg.connection_name}")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Create view for Delta table
        table_name = sanitize_name(Path(cfg.table_path).name) or "delta_table"
        self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM delta_scan('{cfg.table_path}')")

    def _attach_iceberg(self, cfg: SqlConnectionConfig) -> None:
        """Attach Iceberg table."""
        try:
            self._conn.execute("INSTALL iceberg;")
        except Exception:
            pass
        self._conn.execute("LOAD iceberg;")

        alias = _quote_ident(cfg.connection_name)

        if cfg.catalog_type == "rest" and cfg.catalog_uri:
            # REST catalog - attach as Iceberg database
            self._conn.execute(f"ATTACH '{_escape_single_quotes(cfg.catalog_uri)}' AS {alias} (TYPE iceberg);")
        elif cfg.table_path:
            # File-based Iceberg table
            self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")
            table_name = sanitize_name(Path(cfg.table_path).name) or "iceberg_table"
            self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM iceberg_scan('{cfg.table_path}')")
        else:
            raise ValueError(f"iceberg config missing catalog_uri or table_path for {cfg.connection_name}")

    # -------------------------------------------------------------------------
    # Phase 4: Scanner Function Connectors
    # -------------------------------------------------------------------------

    def _attach_odbc(self, cfg: SqlConnectionConfig) -> None:
        """Attach ODBC data source (placeholder schema only)."""
        # Create schema for ODBC connection
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")
        # ODBC tables must be queried explicitly with odbc_scan()
        log.debug("[lazy_attach] ODBC schema created for %s (use odbc_scan for queries)", cfg.connection_name)

    def _attach_gsheets(self, cfg: SqlConnectionConfig) -> None:
        """Attach Google Sheets as a DuckDB view."""
        try:
            self._conn.execute("INSTALL gsheets FROM community;")
        except Exception:
            log.debug("[lazy_attach] gsheets extension not available for %s", cfg.connection_name)
            return

        try:
            self._conn.execute("LOAD gsheets;")
        except Exception as e:
            log.debug("[lazy_attach] Failed loading gsheets: %s", e)
            return

        if not cfg.spreadsheet_id:
            raise ValueError(f"gsheets config missing spreadsheet_id for {cfg.connection_name}")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Configure credentials if provided
        if cfg.credentials_env:
            creds = os.getenv(cfg.credentials_env, "")
            if creds:
                try:
                    self._conn.execute(f"CREATE SECRET gsheet_secret (TYPE gsheet, token='{creds}');")
                except Exception:
                    pass

        # Create view for the sheet
        sheet_ref = cfg.sheet_name or ""
        table_name = (cfg.sheet_name or "sheet").replace(" ", "_").replace("-", "_")

        if sheet_ref:
            self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_gsheet('{cfg.spreadsheet_id}', sheet='{sheet_ref}')")
        else:
            self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_gsheet('{cfg.spreadsheet_id}')")

    def _attach_excel(self, cfg: SqlConnectionConfig) -> None:
        """Attach Excel file as DuckDB views."""
        try:
            self._conn.execute("INSTALL spatial;")
        except Exception:
            pass

        try:
            self._conn.execute("LOAD spatial;")
        except Exception:
            pass

        if not cfg.file_path:
            raise ValueError(f"excel config missing file_path for {cfg.connection_name}")

        file_path = Path(cfg.file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Excel file not found: {cfg.file_path}")

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Discover sheets
        if cfg.sheet_name:
            sheets = [cfg.sheet_name]
        else:
            try:
                import openpyxl
                wb = openpyxl.load_workbook(cfg.file_path, read_only=True)
                sheets = wb.sheetnames
                wb.close()
            except ImportError:
                sheets = ["Sheet1"]
            except Exception:
                sheets = ["Sheet1"]

        for sheet in sheets:
            table_name = sheet.replace(" ", "_").replace("-", "_")
            try:
                self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM st_read('{cfg.file_path}', layer='{sheet}')")
            except Exception:
                try:
                    self._conn.execute(f"CREATE VIEW {alias}.{_quote_ident(table_name)} AS SELECT * FROM read_xlsx('{cfg.file_path}', sheet='{sheet}')")
                except Exception as e:
                    log.debug("[lazy_attach] Failed attaching Excel sheet %s: %s", sheet, e)

    # -------------------------------------------------------------------------
    # Phase 5: Hybrid/Materialization Connectors
    # -------------------------------------------------------------------------

    def _attach_mongodb(self, cfg: SqlConnectionConfig) -> None:
        """Attach MongoDB by materializing collections into DuckDB."""
        try:
            from pymongo import MongoClient
            import pandas as pd
        except ImportError:
            log.warning("[lazy_attach] MongoDB connector requires pymongo and pandas")
            return

        uri = os.getenv(cfg.mongodb_uri_env) if cfg.mongodb_uri_env else None
        if not uri:
            raise ValueError(f"mongodb config missing mongodb_uri_env or env var not set for {cfg.connection_name}")

        if not cfg.database:
            raise ValueError(f"mongodb config missing database for {cfg.connection_name}")

        client = MongoClient(uri, serverSelectionTimeoutMS=5000)
        db = client[cfg.database]

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        for collection_name in db.list_collection_names():
            limit = cfg.sample_row_limit or 1000
            docs = list(db[collection_name].find().limit(limit))

            if not docs:
                continue

            df = pd.json_normalize(docs)
            df.columns = [c.replace(".", "_").replace("$", "") for c in df.columns]

            if "_id" in df.columns:
                df["_id"] = df["_id"].astype(str)

            table_name = collection_name.replace("-", "_").replace(" ", "_")
            temp_name = f'_mongo_{table_name}'

            self._conn.register(temp_name, df)
            self._conn.execute(f"CREATE TABLE {alias}.{_quote_ident(table_name)} AS SELECT * FROM {temp_name}")
            self._conn.unregister(temp_name)

        client.close()

    def _attach_cassandra(self, cfg: SqlConnectionConfig) -> None:
        """Attach Cassandra by materializing tables into DuckDB."""
        try:
            from cassandra.cluster import Cluster
            from cassandra.auth import PlainTextAuthProvider
            import pandas as pd
        except ImportError:
            log.warning("[lazy_attach] Cassandra connector requires cassandra-driver and pandas")
            return

        if not cfg.cassandra_hosts:
            raise ValueError(f"cassandra config missing cassandra_hosts for {cfg.connection_name}")

        if not cfg.cassandra_keyspace:
            raise ValueError(f"cassandra config missing cassandra_keyspace for {cfg.connection_name}")

        auth = None
        if cfg.user:
            password = os.getenv(cfg.password_env) if cfg.password_env else None
            auth = PlainTextAuthProvider(cfg.user, password or "")

        cluster = Cluster(cfg.cassandra_hosts, auth_provider=auth)
        session = cluster.connect(cfg.cassandra_keyspace)

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        tables_query = f"SELECT table_name FROM system_schema.tables WHERE keyspace_name = '{cfg.cassandra_keyspace}'"
        tables = session.execute(tables_query)

        for row in tables:
            table_name = row.table_name
            limit = cfg.sample_row_limit or 1000

            try:
                sample_query = f"SELECT * FROM {table_name} LIMIT {limit}"
                rows = session.execute(sample_query)

                df = pd.DataFrame(list(rows))
                if df.empty:
                    continue

                temp_name = f'_cass_{table_name}'
                self._conn.register(temp_name, df)
                self._conn.execute(f"CREATE TABLE {alias}.{_quote_ident(table_name)} AS SELECT * FROM {temp_name}")
                self._conn.unregister(temp_name)
            except Exception as e:
                log.debug("[lazy_attach] Failed materializing Cassandra table %s: %s", table_name, e)

        cluster.shutdown()

    def _attach_clickhouse(self, cfg: SqlConnectionConfig) -> None:
        """Attach ClickHouse by materializing tables into DuckDB."""
        try:
            import clickhouse_connect
            import pandas as pd
        except ImportError:
            log.warning("[lazy_attach] ClickHouse connector requires clickhouse-connect and pandas: pip install clickhouse-connect")
            return

        if not cfg.host:
            raise ValueError(f"clickhouse config missing host for {cfg.connection_name}")

        # Get connection parameters
        host = cfg.host
        port = cfg.port or 8123  # ClickHouse HTTP port (default)
        database = cfg.database or "default"
        user = cfg.user or "default"
        password = cfg.password or (os.getenv(cfg.password_env) if cfg.password_env else "")

        try:
            client = clickhouse_connect.get_client(
                host=host,
                port=port,
                database=database,
                username=user,
                password=password,
            )
        except Exception as e:
            log.warning("[lazy_attach] Failed to connect to ClickHouse %s:%s: %s", host, port, e)
            return

        # Create schema
        alias = _quote_ident(cfg.connection_name)
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {alias};")

        # Get list of tables in the database
        try:
            tables_result = client.query(f"SHOW TABLES FROM {database}")
            table_names = [row[0] for row in tables_result.result_rows]
        except Exception as e:
            log.warning("[lazy_attach] Failed to list ClickHouse tables: %s", e)
            client.close()
            return

        materialized_count = 0
        total_rows = 0

        for table_name in table_names:
            limit = cfg.sample_row_limit or 1000

            try:
                # Get sample data
                df = client.query_df(f"SELECT * FROM {database}.{table_name} LIMIT {limit}")

                if df.empty:
                    continue

                # Sanitize table name for DuckDB
                safe_table_name = table_name.replace("-", "_").replace(" ", "_")
                temp_name = f'_ch_{safe_table_name}'

                self._conn.register(temp_name, df)
                self._conn.execute(f"CREATE TABLE {alias}.{_quote_ident(safe_table_name)} AS SELECT * FROM {temp_name}")
                self._conn.unregister(temp_name)

                materialized_count += 1
                total_rows += len(df)
                print(f"    ✓ Materialized ClickHouse table: {table_name} → {cfg.connection_name}.{safe_table_name} ({len(df)} rows)")

            except Exception as e:
                log.debug("[lazy_attach] Failed materializing ClickHouse table %s: %s", table_name, e)
                print(f"    ⚠️  Failed to materialize {table_name}: {str(e)[:60]}")

        client.close()
        print(f"  └─ Materialized ClickHouse: {database} ({materialized_count} tables, {total_rows} rows)")

    def _attach_duckdb_file_if_present(self, db_name: str) -> bool:
        # Avoid reserved/system catalogs
        if db_name.lower() in {"system", "temp", "memory"}:
            return False

        for folder_cfg in self._duckdb_folder_configs:
            folder_path = folder_cfg.folder_path
            if not folder_path:
                continue

            file_map = self._duckdb_file_maps.get(folder_path)
            if file_map is None:
                file_map = self._build_duckdb_file_map(folder_path)
                self._duckdb_file_maps[folder_path] = file_map

            db_file = file_map.get(db_name)
            if not db_file:
                continue

            _attach_duckdb_read_only_with_snapshot_fallback(self._conn, db_file, db_name)
            return True

        return False

    def _build_duckdb_file_map(self, folder_path: str) -> Dict[str, Path]:
        folder = Path(folder_path)
        if not folder.exists() or not folder.is_dir():
            return {}
        mapping: Dict[str, Path] = {}
        for db_file in folder.glob("*.duckdb"):
            mapping[sanitize_name(db_file.name)] = db_file
        return mapping

    def _ensure_csv_tables(self, connection_name: str, tables: Set[str]) -> None:
        cfg = self._configs.get(connection_name)
        if not cfg or cfg.type != "csv_folder":
            return
        if not cfg.folder_path:
            return

        # Ensure schema exists
        self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {_quote_ident(connection_name)};")

        # Build file map once per connection
        file_map = self._csv_file_maps.get(connection_name)
        if file_map is None:
            file_map = self._build_csv_file_map(cfg.folder_path)
            self._csv_file_maps[connection_name] = file_map

        for table_name in sorted(tables):
            try:
                if self._csv_table_exists(connection_name, table_name):
                    continue

                csv_file = file_map.get(table_name)
                if not csv_file:
                    continue

                full_name = f"{_quote_ident(connection_name)}.{_quote_ident(table_name)}"

                if _lazy_attach_csv_materialize():
                    self._conn.execute(
                        f"""
                        CREATE TABLE {full_name} AS
                        SELECT * FROM read_csv_auto('{_escape_single_quotes(str(csv_file))}', AUTO_DETECT=TRUE, ignore_errors=true)
                        """.strip()
                    )
                else:
                    self._conn.execute(
                        f"""
                        CREATE OR REPLACE VIEW {full_name} AS
                        SELECT * FROM read_csv_auto('{_escape_single_quotes(str(csv_file))}', AUTO_DETECT=TRUE, ignore_errors=true)
                        """.strip()
                    )
            except Exception as e:
                log.debug("[lazy_attach] Failed ensuring csv table %s.%s: %s", connection_name, table_name, e)

    def _csv_table_exists(self, schema_name: str, table_name: str) -> bool:
        try:
            row = self._conn.execute(
                """
                SELECT COUNT(*) FROM information_schema.tables
                WHERE table_schema = ? AND table_name = ?
                """,
                [schema_name, table_name],
            ).fetchone()
            return bool(row and row[0] > 0)
        except Exception:
            return False

    def _build_csv_file_map(self, folder_path: str) -> Dict[str, Path]:
        folder = Path(folder_path)
        if not folder.exists() or not folder.is_dir():
            return {}
        mapping: Dict[str, Path] = {}
        for csv_file in folder.glob("*.csv"):
            mapping[sanitize_name(csv_file.name)] = csv_file
        return mapping

    # ---------------------------------------------------------------------
    # ClickHouse table materialization
    # ---------------------------------------------------------------------

    def _ensure_clickhouse_tables(self, connection_name: str, tables: Set[str]) -> None:
        """
        Materialize ClickHouse tables into DuckDB on demand.

        Since DuckDB doesn't have native ClickHouse support, we fetch
        the table data via the ClickHouse Python client and register
        it as a DuckDB table.
        """
        cfg = self._configs.get(connection_name)
        if not cfg or cfg.type != "clickhouse":
            return

        # Ensure schema exists
        try:
            self._conn.execute(f"CREATE SCHEMA IF NOT EXISTS {_quote_ident(connection_name)};")
        except Exception as e:
            log.debug("[lazy_attach] Failed creating schema %s: %s", connection_name, e)
            return

        # Track which tables we've already materialized
        if connection_name not in self._clickhouse_tables:
            self._clickhouse_tables[connection_name] = set()

        already_done = self._clickhouse_tables[connection_name]

        for table_name in sorted(tables):
            if table_name in already_done:
                continue

            try:
                # Check if already exists in DuckDB
                if self._csv_table_exists(connection_name, table_name):
                    already_done.add(table_name)
                    continue

                # Fetch from ClickHouse and materialize
                self._materialize_clickhouse_table(cfg, connection_name, table_name)
                already_done.add(table_name)

            except Exception as e:
                log.warning(
                    "[lazy_attach] Failed materializing ClickHouse table %s.%s: %s",
                    connection_name, table_name, e
                )

    def _materialize_clickhouse_table(
        self,
        cfg: SqlConnectionConfig,
        schema_name: str,
        table_name: str,
    ) -> None:
        """
        Fetch a table from ClickHouse and create it in DuckDB.
        """
        try:
            from rvbbit.db_adapter import get_db
        except ImportError:
            log.warning("[lazy_attach] ClickHouse db_adapter not available")
            return

        db = get_db()
        if not db:
            log.warning("[lazy_attach] No ClickHouse connection available")
            return

        # Determine the ClickHouse table name
        # If cfg.database is set, use it; otherwise use connection_name
        ch_database = cfg.database or schema_name
        ch_table = f"{ch_database}.{table_name}"

        log.debug("[lazy_attach] Materializing ClickHouse table %s -> %s.%s", ch_table, schema_name, table_name)

        try:
            # Fetch data from ClickHouse
            rows = db.query(f"SELECT * FROM {ch_table}")

            if not rows:
                # Create empty table - get schema from ClickHouse
                schema_rows = db.query(f"DESCRIBE TABLE {ch_table}")
                if schema_rows:
                    # Create empty table with correct schema
                    columns = []
                    for row in schema_rows:
                        col_name = row.get('name', row.get('column_name', ''))
                        col_type = row.get('type', 'VARCHAR')
                        # Map ClickHouse types to DuckDB types
                        duckdb_type = _clickhouse_type_to_duckdb(col_type)
                        if col_name:
                            columns.append(f"{_quote_ident(col_name)} {duckdb_type}")

                    if columns:
                        full_name = f"{_quote_ident(schema_name)}.{_quote_ident(table_name)}"
                        create_sql = f"CREATE TABLE {full_name} ({', '.join(columns)})"
                        self._conn.execute(create_sql)
                        log.debug("[lazy_attach] Created empty table %s", full_name)
                return

            # Convert to DataFrame and register
            import pandas as pd
            df = pd.DataFrame(rows)

            full_name = f"{_quote_ident(schema_name)}.{_quote_ident(table_name)}"

            # Register DataFrame as a DuckDB table
            self._conn.register('_ch_temp_df', df)
            self._conn.execute(f"CREATE TABLE {full_name} AS SELECT * FROM _ch_temp_df")
            self._conn.unregister('_ch_temp_df')

            log.debug("[lazy_attach] Materialized ClickHouse table %s (%d rows)", full_name, len(df))

        except Exception as e:
            log.error("[lazy_attach] Error fetching from ClickHouse %s: %s", ch_table, e)
            raise


def _clickhouse_type_to_duckdb(ch_type: str) -> str:
    """Map ClickHouse types to DuckDB types."""
    ch_type_upper = ch_type.upper()

    # Handle Nullable wrapper
    if ch_type_upper.startswith('NULLABLE('):
        inner = ch_type[9:-1]  # Strip Nullable(...)
        return _clickhouse_type_to_duckdb(inner)

    # Basic type mappings
    type_map = {
        'STRING': 'VARCHAR',
        'FIXEDSTRING': 'VARCHAR',
        'UUID': 'VARCHAR',
        'INT8': 'TINYINT',
        'INT16': 'SMALLINT',
        'INT32': 'INTEGER',
        'INT64': 'BIGINT',
        'UINT8': 'UTINYINT',
        'UINT16': 'USMALLINT',
        'UINT32': 'UINTEGER',
        'UINT64': 'UBIGINT',
        'FLOAT32': 'FLOAT',
        'FLOAT64': 'DOUBLE',
        'BOOL': 'BOOLEAN',
        'BOOLEAN': 'BOOLEAN',
        'DATE': 'DATE',
        'DATE32': 'DATE',
        'DATETIME': 'TIMESTAMP',
        'DATETIME64': 'TIMESTAMP',
    }

    # Check exact matches
    for ch, duck in type_map.items():
        if ch_type_upper == ch or ch_type_upper.startswith(ch + '('):
            return duck

    # Handle Decimal
    if ch_type_upper.startswith('DECIMAL'):
        return ch_type  # DuckDB supports DECIMAL(p,s) syntax

    # Handle Array types
    if ch_type_upper.startswith('ARRAY('):
        inner = ch_type[6:-1]
        inner_duck = _clickhouse_type_to_duckdb(inner)
        return f"{inner_duck}[]"

    # Handle Enum
    if ch_type_upper.startswith('ENUM'):
        return 'VARCHAR'

    # Default fallback
    return 'VARCHAR'


def _lazy_attach_enabled() -> bool:
    val = os.environ.get("RVBBIT_LAZY_ATTACH", "1").strip().lower()
    return val not in ("0", "false", "no", "off")


def _lazy_attach_csv_materialize() -> bool:
    val = os.environ.get("RVBBIT_LAZY_ATTACH_CSV_MATERIALIZE", "0").strip().lower()
    return val in ("1", "true", "yes", "on")


def _attach_duckdb_read_only_with_snapshot_fallback(duckdb_conn, db_file: Path, db_name: str, max_retries: int = 2) -> None:
    """
    Attach a DuckDB file READ_ONLY with snapshot fallback (copy-on-read) if locked.

    Adapted from `DatabaseConnector._attach_duckdb_file`, but scoped for the
    pgwire server's per-session DuckDB connection.
    """
    last_error = None

    # Strategy 1/2: direct attach with retries
    for attempt in range(max_retries):
        try:
            duckdb_conn.execute(f"ATTACH '{_escape_single_quotes(str(db_file))}' AS {_quote_ident(db_name)} (READ_ONLY)")
            return
        except Exception as e:
            last_error = e
            error_str = str(e).lower()
            is_lock_error = any(
                phrase in error_str
                for phrase in ("lock", "could not set lock", "database is locked", "unable to open", "exclusive")
            )
            if not is_lock_error:
                raise
            if attempt < max_retries - 1:
                time.sleep(0.3 * (2**attempt))

    # Strategy 3: snapshot copy
    temp_dir = Path(tempfile.gettempdir()) / "rvbbit_duckdb_snapshots"
    temp_dir.mkdir(exist_ok=True)
    temp_path = temp_dir / f"{db_name}.duckdb"

    try:
        shutil.copy2(db_file, temp_path)
        duckdb_conn.execute(f"ATTACH '{_escape_single_quotes(str(temp_path))}' AS {_quote_ident(db_name)} (READ_ONLY)")
        return
    except Exception as copy_error:
        raise Exception(
            f"Failed to attach {db_file.name}: direct attach failed ({last_error}), snapshot copy failed ({copy_error})"
        )
